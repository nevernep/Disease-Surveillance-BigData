"""Export the PostgreSQL warehouse to one Excel workbook for Power BI on the web.

Power BI Service (app.powerbi.com) cannot reach a database on a laptop without
an on-premises gateway, but it can import an uploaded Excel file. Each table
is an Excel Table named after the warehouse table (e.g. ``dim_sex``) on a sheet
named ``sheet_<table>``. In Power BI pick the names WITHOUT the ``sheet_``
prefix: Excel Tables always carry their header row, whereas a sheet of
text-only columns (dim_sex, data_quality) is imported as Column1, Column2.

Writes data/processed/powerbi/disease_surveillance_powerbi.xlsx.
Does not import pyspark.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Dict

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from spark.load_warehouse import connection_settings

# Star schema tables, in the order they appear in the workbook.
TABLES: Dict[str, str] = {
    "fact_disease_cases": "SELECT * FROM mart.fact_disease_cases ORDER BY 1, 2, 3, 4, 5",
    "fact_population": "SELECT * FROM mart.fact_population ORDER BY 1, 2",
    "dim_date": "SELECT * FROM mart.dim_date ORDER BY date_key",
    "dim_district": "SELECT * FROM mart.dim_district ORDER BY district_key",
    "dim_disease": "SELECT * FROM mart.dim_disease ORDER BY disease_key",
    "dim_age_group": "SELECT * FROM mart.dim_age_group ORDER BY age_group_key",
    "dim_sex": "SELECT * FROM mart.dim_sex ORDER BY sex",
}


def read_tables() -> Dict[str, pd.DataFrame]:
    import psycopg2

    frames: Dict[str, pd.DataFrame] = {}
    connection = psycopg2.connect(**connection_settings())
    try:
        with connection.cursor() as cursor:
            for name, sql in TABLES.items():
                cursor.execute(sql)
                columns = [description[0] for description in cursor.description]
                frames[name] = pd.DataFrame(cursor.fetchall(), columns=columns)
    finally:
        connection.close()

    # Real dates so Power BI can use dim_date as a date table.
    frames["dim_date"]["month_start"] = pd.to_datetime(
        frames["dim_date"]["month_start"]
    )
    return frames


def write_workbook(frames: Dict[str, pd.DataFrame], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".tmp.xlsx")

    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        for name, frame in frames.items():
            sheet_name = f"sheet_{name}"[:31]
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.sheets[sheet_name]
            last_cell = f"{get_column_letter(len(frame.columns))}{len(frame) + 1}"
            table = Table(displayName=name, ref=f"A1:{last_cell}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True
            )
            sheet.add_table(table)
            for index, column in enumerate(frame.columns, start=1):
                sheet.column_dimensions[get_column_letter(index)].width = max(
                    12, len(str(column)) + 2
                )

    temporary.replace(output)


def add_quality_report(frames: Dict[str, pd.DataFrame], project_root: Path) -> None:
    report = project_root / "data" / "processed" / "quality" / "data_quality_report.csv"
    if report.is_file():
        frames["data_quality"] = pd.read_csv(report, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Export the warehouse to an Excel workbook for Power BI web"
    )
    parser.add_argument("--project-root", type=Path, default=Path("."))
    args = parser.parse_args()
    project_root = args.project_root.resolve()

    frames = read_tables()
    add_quality_report(frames, project_root)
    output = (
        project_root / "data" / "processed" / "powerbi"
        / "disease_surveillance_powerbi.xlsx"
    )
    write_workbook(frames, output)

    for name, frame in frames.items():
        print(f"{name:<20} {len(frame):>7} rows")
    print(f"Wrote {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
